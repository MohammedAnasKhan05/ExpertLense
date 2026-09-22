"""
Tests for multi-format document text extractor (.txt, .pdf, .docx, .doc).
"""

import io
import pytest
from app.services.document_extractor import (
    extract_text_from_file,
    _extract_txt,
    _extract_pdf,
    _extract_docx,
)


def test_extract_txt_utf8():
    content = "Expert 1 – Dr. Emily Carter\nMarket: United Kingdom\n\n00:00\nInterviewer: Hello".encode("utf-8")
    result = extract_text_from_file("transcript.txt", content)
    assert "Dr. Emily Carter" in result
    assert "United Kingdom" in result


def test_extract_txt_latin1():
    content = "Expert - Dr. Jean Martin\nMarket: France\n\n00:00\nInterviewer: Bonjour".encode("latin-1")
    result = extract_text_from_file("transcript.txt", content)
    assert "Dr. Jean Martin" in result


def test_extract_docx():
    import docx
    doc = docx.Document()
    doc.add_paragraph("Expert 2 – Anna Keller")
    doc.add_paragraph("Market: Germany")
    doc.add_paragraph("00:00")
    doc.add_paragraph("Interviewer: Guten Tag")

    bio = io.BytesIO()
    doc.save(bio)
    content = bio.getvalue()

    result = extract_text_from_file("test_transcript.docx", content)
    assert "Anna Keller" in result
    assert "Germany" in result
    assert "Guten Tag" in result


def test_extract_pdf():
    import pypdf
    writer = pypdf.PdfWriter()
    # Create page with text if possible, or verify empty handling
    # pypdf Writer doesn't have an easy text drawer without reportlab, but let's test invalid/empty PDF gracefully raises ValueError
    bio = io.BytesIO()
    writer.write(bio)
    empty_pdf_bytes = bio.getvalue()

    with pytest.raises(ValueError, match="No extractable text found in PDF"):
        extract_text_from_file("test.pdf", empty_pdf_bytes)


def test_unsupported_format():
    with pytest.raises(ValueError, match="Unsupported file format"):
        extract_text_from_file("test.xyz", b"some binary data")
